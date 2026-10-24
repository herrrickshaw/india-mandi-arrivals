"""Crop ranking workbook: state-wise mandi arrivals and price volatility for 23 crops (Agmarknet).

Usage: python analysis/workbooks/crop_ranking.py <out.xlsx> <results.json>
Reads data/agmarknet/state_monthly.csv.gz for the 23 field crops in crosswalks/commodity_crosswalk.csv.
Wide sheets (one column per crop-state pair) hold prices and arrivals as data; Returns and Ranking are formulas.
results.json carries the Python cross-check of the same metrics for charting.
"""
import csv, gzip, pathlib, sys, json, statistics, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter as L

ROOT = pathlib.Path(__file__).resolve().parents[2]
OUT, RES = sys.argv[1:3]
CORE = {r["agmarknet_commodity"] for r in csv.DictReader(open(ROOT / "crosswalks/commodity_crosswalk.csv"))}
F = "Arial"; BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); BOLD = Font(name=F, bold=True)
H1 = Font(name=F, bold=True, size=13); HDR = Font(name=F, bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F4E78")
RS = '#,##0;(#,##0);-'; PCT = '0.0%;(0.0%);-'
MIN_T, START = 50, "2023-01"

data = collections.defaultdict(dict)   # (crop, state) -> month -> (price, arrivals)
for r in csv.DictReader(gzip.open(ROOT / "data/agmarknet/state_monthly.csv.gz", "rt", encoding="utf-8")):
    a = float(r["arrivals_t"] or 0)
    if r["commodity"] in CORE and r["price_rs_qtl"] and a >= MIN_T:
        data[(r["commodity"], r["state"])][r["month"]] = (float(r["price_rs_qtl"]), a)
months = sorted({m for v in data.values() for m in v}); N = len(months)
last12 = months[-12:]; rmonths = [m for m in months if m >= START]

# All-India per crop: arrival-weighted mean price, summed arrivals
crops = sorted({c for c, _ in data})
for c in crops:
    ai = {}
    for m in months:
        pts = [v[m] for (cc, st), v in data.items() if cc == c and st != "All India" and m in v]
        w = sum(a for _, a in pts)
        if w: ai[m] = (sum(p * a for p, a in pts) / w, w)
    data[(c, "All India")] = ai

def metrics(v):
    rets = [abs(v[m][0] / v[p][0] - 1) for p, m in zip(months, months[1:]) if m >= START and m in v and p in v]
    prices = [v[m][0] for m in rmonths if m in v]
    qs = statistics.quantiles(prices, n=20, method="inclusive") if len(prices) >= 24 else None
    return dict(arr12=sum(v[m][1] for m in last12 if m in v), vol=statistics.median(rets) if len(rets) >= 24 else None,
                n=len(rets), rng=(qs[-1] / qs[0] - 1) if qs else None,
                last=v[max(v)][0] if v else None)
res = {f"{c}|{s}": metrics(v) for (c, s), v in data.items()}
# keep pairs with enough history and real volume
pairs = sorted([k for k, m in res.items() if m["vol"] is not None and (k.endswith("|All India") or m["arr12"] >= 5000)],
               key=lambda k: (k.split("|")[0], -res[k]["arr12"]))
json.dump({"pairs": pairs, "res": res, "months": months}, open(RES, "w"))

wb = Workbook()
def hdr(ws, r, vals, c0=1):
    for i, v in enumerate(vals):
        c = ws.cell(r, c0 + i, v); c.font = HDR; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical="center")

ws = wb.active; ws.title = "Notes"
ws["A1"] = "Crop ranking by state: mandi arrivals and price volatility, 23 crops, Agmarknet Jan 2019 to Oct 2026"; ws["A1"].font = H1
notes = [
 "Source: Agmarknet dashboard API (api.agmarknet.gov.in/v1/dashboard-data, statewise_price_arrival_sp). Monthly weighted-average price (Rs/quintal) and arrivals (tonnes) by state.",
 f"Volatility = median absolute month-on-month price change (the typical monthly move, up or down), months from {START} to Oct 2026 where both months are present. The median is used so that a single misreported month cannot dominate. A state-month with under {MIN_T} t of arrivals is dropped as unreliable.",
 "Price range = 95th-percentile monthly price / 5th-percentile monthly price - 1, same window (trims the most extreme months). Arrivals = total over the last 12 months (Nov 2025 to Oct 2026).",
 "Pairs are kept with 24+ monthly price changes and 5,000+ t of arrivals in the last 12 months. All-India = arrival-weighted average of state prices.",
 "Caveats: Agmarknet widened its market coverage from Nov 2025, so recent arrivals are higher than earlier years for every crop. Prices mix varieties and grades within a crop. Oct 2026 runs to 23 Oct.",
]
for i, t in enumerate(notes):
    c = ws.cell(3 + i, 1, t); c.font = BLACK; c.alignment = Alignment(wrap_text=True)
ws.column_dimensions["A"].width = 130
ws.cell(9, 1, "Volatility window starts (edit to change):").font = BOLD
c = ws.cell(9, 2, START); c.font = BLUE
ws.column_dimensions["B"].width = 12

R0 = 4; RN = R0 + N - 1
def wide(title, idx, desc):
    s = wb.create_sheet(title); s["A1"] = desc; s["A1"].font = H1
    s.cell(2, 1, "Crop").font = BOLD; s.cell(3, 1, "State / Month").font = BOLD
    for j, k in enumerate(pairs):
        c, st = k.split("|")
        s.cell(2, 2 + j, c).font = BOLD; h = s.cell(3, 2 + j, st); h.font = HDR; h.fill = HFILL
        v = data[(c, st)]
        for i, m in enumerate(months):
            if m in v:
                x = s.cell(R0 + i, 2 + j, round(v[m][idx], 1)); x.font = BLUE; x.number_format = RS
    for i, m in enumerate(months): s.cell(R0 + i, 1, m).font = BLACK
    s.freeze_panes = "B4"; return s
P = wide("Prices", 0, "Monthly weighted-average price (Rs/quintal) by crop and state")
A = wide("Arrivals", 1, "Monthly arrivals (tonnes) by crop and state")

Rt = wb.create_sheet("Returns"); Rt["A1"] = "Absolute month-on-month price change (formula: ABS(this month / previous month - 1); blank if either is missing)"; Rt["A1"].font = H1
for j, k in enumerate(pairs):
    c, st = k.split("|"); col = L(2 + j)
    Rt.cell(2, 2 + j, c).font = BOLD; h = Rt.cell(3, 2 + j, st); h.font = HDR; h.fill = HFILL
    for i in range(1, N):
        r = R0 + i
        Rt.cell(r, 2 + j, f"=IF(OR(Prices!{col}{r}=\"\",Prices!{col}{r-1}=\"\"),\"\",ABS(Prices!{col}{r}/Prices!{col}{r-1}-1))").number_format = PCT
for i, m in enumerate(months): Rt.cell(R0 + i, 1, f"=Prices!A{R0 + i}").font = BLACK
Rt.freeze_panes = "B4"

# Ranking sheet: one row per pair; window rows located by MATCH on the start month
rk = wb.create_sheet("Ranking", 1)
rk["A1"] = "Ranking: arrivals (last 12 months) and price volatility by crop and state"; rk["A1"].font = H1
rk["A2"] = "Ranks: 1 = largest arrivals / most volatile, within the crop (state ranks) and across crops (All-India rows). Green = steadier, red = more volatile."
hdr(rk, 3, ["Crop", "State", "Arrivals, last 12m (t)", "Share of crop's arrivals", "Volatility (median monthly move)", "Monthly changes counted",
            "Price range (p95/p5 - 1)", "Latest price (Rs/qtl)", "Rank in crop by arrivals", "Rank in crop by volatility", "All-India rank by arrivals", "All-India rank by volatility"])
row = 4; first = {}
startrow = f"(MATCH(Notes!$B$9,Prices!$A${R0}:$A${RN},0)+{R0 - 1})"
for j, k in enumerate(pairs):
    c, st = k.split("|"); col = L(2 + j)
    first.setdefault(c, row)
    rk.cell(row, 1, c).font = BLACK; rk.cell(row, 2, st).font = BOLD if st == "All India" else BLACK
    rk.cell(row, 3, f"=SUM(Arrivals!{col}{RN-11}:{col}{RN})").number_format = RS
    rng = lambda sh: f"INDIRECT(\"{sh}!{col}\"&{startrow}&\":{col}{RN}\")"
    rk.cell(row, 5, f"=IFERROR(MEDIAN({rng('Returns')}),\"\")").number_format = PCT
    rk.cell(row, 6, f"=COUNT({rng('Returns')})")
    rk.cell(row, 7, f"=IFERROR(PERCENTILE({rng('Prices')},0.95)/PERCENTILE({rng('Prices')},0.05)-1,\"\")").number_format = PCT
    rk.cell(row, 8, f"=INDEX(Prices!{col}${R0}:{col}${RN},MATCH(9.99E+307,Prices!{col}${R0}:{col}${RN}))").number_format = RS
    row += 1
last = row - 1
for r in range(4, row):
    c = rk.cell(r, 1).value; st = rk.cell(r, 2).value
    a0 = first[c]; a1 = max(rr for rr in range(4, row) if rk.cell(rr, 1).value == c)
    if st != "All India":
        rk.cell(r, 4, f"=IFERROR(C{r}/C{a0},\"\")").number_format = PCT
        rk.cell(r, 9, f"=COUNTIFS($A${a0}:$A${a1},A{r},$B${a0}:$B${a1},\"<>All India\",$C${a0}:$C${a1},\">\"&C{r})+1")
        rk.cell(r, 10, f"=IF(E{r}=\"\",\"\",COUNTIFS($A${a0}:$A${a1},A{r},$B${a0}:$B${a1},\"<>All India\",$E${a0}:$E${a1},\">\"&E{r})+1)")
    else:
        rk.cell(r, 11, f"=COUNTIFS($B$4:$B${last},\"All India\",$C$4:$C${last},\">\"&C{r})+1")
        rk.cell(r, 12, f"=IF(E{r}=\"\",\"\",COUNTIFS($B$4:$B${last},\"All India\",$E$4:$E${last},\">\"&E{r})+1)")
rk.conditional_formatting.add(f"E4:E{last}", ColorScaleRule(start_type="percentile", start_value=5, start_color="63BE7B", mid_type="percentile", mid_value=50, mid_color="FFFFFF", end_type="percentile", end_value=95, end_color="F8696B"))
for cidx, w in zip(range(1, 13), (24, 20, 14, 11, 13, 10, 12, 11, 10, 10, 10, 10)): rk.column_dimensions[L(cidx)].width = w
rk.row_dimensions[3].height = 48; rk.freeze_panes = "C4"

for wsx in wb.worksheets:
    for rr in wsx.iter_rows():
        for cell in rr:
            if cell.value is not None and cell.font.name != F:
                cell.font = Font(name=F, bold=cell.font.bold, color=cell.font.color, italic=cell.font.italic, size=cell.font.size)
wb.save(OUT); print("saved", OUT, len(pairs), "pairs,", len(crops), "crops")
